from pathlib import Path

import docx
import pdfplumber


def extract_text(file_path: str) -> str:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".docx":
        return _extract_docx(file_path)
    if suffix == ".pdf":
        return _extract_pdf(file_path)
    if suffix in (".md", ".markdown", ".txt"):
        return Path(file_path).read_text(encoding="utf-8")
    raise ValueError(f"Unsupported resume file type: {suffix}")


def _extract_docx(file_path: str) -> str:
    document = docx.Document(file_path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pdf(file_path: str) -> str:
    parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
    return "\n".join(parts)
